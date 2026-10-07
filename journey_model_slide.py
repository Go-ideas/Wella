from __future__ import annotations

import io
from typing import Final


TITLE: Final[str] = "Cómo leer el Journey de decisión de compra"
SUBTITLE: Final[str] = "Modelo de Red Bayesiana Secuencial"


def journey_slide_html() -> str:
    """Return a dashboard-ready explanatory slide as self-contained HTML."""
    return """
    <style>
      .journey-slide-wrap{
        border:1px solid #DCE7F0;
        border-radius:28px;
        padding:26px;
        background:linear-gradient(135deg,#F7FAFD 0%,#FFFFFF 48%,#F2F7FC 100%);
        box-shadow:0 14px 38px rgba(18,52,86,.08);
        color:#143A5A;
        font-family:Inter,Segoe UI,Arial,sans-serif;
      }
      .journey-slide-kicker{font-size:.78rem;font-weight:800;letter-spacing:.11em;text-transform:uppercase;color:#6F86A0;margin-bottom:6px}
      .journey-slide-title{font-size:2.05rem;font-weight:900;line-height:1.08;margin:0;color:#10385E}
      .journey-slide-subtitle{font-size:1.05rem;color:#6E8399;margin-top:8px;max-width:760px}
      .journey-model-band{display:grid;grid-template-columns:1fr 84px 1fr 84px 1fr;gap:18px;align-items:center;margin:28px 0 24px}
      .journey-step{min-height:178px;border:1.5px solid #D8E5EF;border-radius:24px;background:#FFFFFF;padding:20px 20px 18px;box-shadow:0 10px 28px rgba(27,74,112,.07)}
      .journey-step.primary{border-color:#2F80ED;background:#F3F8FF}
      .journey-step-number{width:42px;height:42px;border-radius:50%;background:#2F80ED;color:white;display:flex;align-items:center;justify-content:center;font-weight:900;font-size:1.12rem;margin-bottom:13px}
      .journey-step h3{font-size:1.32rem;margin:0 0 8px;color:#153C61}
      .journey-step p{font-size:.97rem;line-height:1.42;color:#6D8399;margin:0}
      .journey-arrow{height:4px;background:#94B6D5;border-radius:99px;position:relative}
      .journey-arrow:after{content:"";position:absolute;right:-2px;top:-6px;width:15px;height:15px;border-top:4px solid #94B6D5;border-right:4px solid #94B6D5;transform:rotate(45deg)}
      .journey-explain-grid{display:grid;grid-template-columns:1.05fr 1fr 1fr;gap:16px;margin-top:8px}
      .journey-info-card{border:1px solid #DCE7F0;border-radius:20px;background:rgba(255,255,255,.78);padding:18px 18px 16px}
      .journey-info-card h4{margin:0 0 8px;font-size:.86rem;letter-spacing:.09em;text-transform:uppercase;color:#7690A9;font-weight:900}
      .journey-info-card p{margin:0;color:#163E63;font-size:1.02rem;line-height:1.45;font-weight:650}
      .journey-formula{font-family:Consolas,Menlo,monospace;background:#0E3A63;color:white;border-radius:16px;padding:14px 16px;font-weight:800;margin-top:10px;font-size:.93rem}
      .journey-footer{margin-top:20px;border-top:1px solid #DCE7F0;padding-top:14px;display:flex;gap:12px;align-items:center;color:#6B839B;font-size:.9rem}
      .journey-footer-badge{background:#EAF3FB;color:#1D6FB5;border:1px solid #CFE4F7;border-radius:999px;padding:7px 12px;font-weight:850;white-space:nowrap}
      @media(max-width:900px){.journey-model-band{grid-template-columns:1fr}.journey-arrow{height:30px;width:4px;margin:auto}.journey-arrow:after{right:-6px;top:18px;transform:rotate(135deg)}.journey-explain-grid{grid-template-columns:1fr}.journey-slide-title{font-size:1.55rem}}
    </style>
    <div class="journey-slide-wrap">
      <div class="journey-slide-kicker">Decisión de compra · Modelo analítico</div>
      <h2 class="journey-slide-title">Cómo leer el Journey de decisión de compra</h2>
      <div class="journey-slide-subtitle">El modelo estima cómo avanza la decisión: qué criterio aparece primero, cuál entra después y qué termina cerrando la compra.</div>

      <div class="journey-model-band">
        <div class="journey-step primary">
          <div class="journey-step-number">1</div>
          <h3>Primero</h3>
          <p>Identifica el criterio que abre la decisión de compra. Ejemplo: tono/color.</p>
        </div>
        <div class="journey-arrow"></div>
        <div class="journey-step">
          <div class="journey-step-number">2</div>
          <h3>Después</h3>
          <p>Estima qué criterio aparece después, condicionado por el primer paso. Ejemplo: precio o beneficios.</p>
        </div>
        <div class="journey-arrow"></div>
        <div class="journey-step">
          <div class="journey-step-number">3</div>
          <h3>Cierre</h3>
          <p>Muestra qué termina definiendo la compra, dado el camino anterior. Ejemplo: disponibilidad.</p>
        </div>
      </div>

      <div class="journey-explain-grid">
        <div class="journey-info-card">
          <h4>Qué significan los %</h4>
          <p>Son probabilidades acumuladas desde el total. Por eso bajan conforme avanza el Journey.</p>
          <div class="journey-formula">P(ruta) = P(D1) × P(D2|D1) × P(D3|D1,D2)</div>
        </div>
        <div class="journey-info-card">
          <h4>Certeza</h4>
          <p>La estabilidad combina simulaciones posteriores, intervalo creíble y base efectiva. Si la base es baja, se interpreta como exploratorio.</p>
        </div>
        <div class="journey-info-card">
          <h4>Para qué sirve</h4>
          <p>Ayuda a decidir cómo ordenar el anaquel, qué comunicar primero y qué reforzar para cerrar la compra.</p>
        </div>
      </div>

      <div class="journey-footer">
        <span class="journey-footer-badge">Lectura ejecutiva</span>
        <span>El cliente ve un Journey simple; el modelo calcula probabilidades bayesianas detrás.</span>
      </div>
    </div>
    """


def build_journey_model_pptx() -> bytes:
    """Build a one-slide PPTX explaining the Sequential Bayesian Journey."""
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except Exception as exc:  # pragma: no cover - UI fallback
        raise RuntimeError("python-pptx is required to export the explanatory slide") from exc

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    navy = RGBColor(16, 56, 94)
    blue = RGBColor(47, 128, 237)
    mid = RGBColor(111, 134, 160)
    pale = RGBColor(243, 248, 255)
    line = RGBColor(216, 229, 239)
    white = RGBColor(255, 255, 255)

    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.28), Inches(0.25), Inches(12.77), Inches(7.0))
    bg.fill.solid(); bg.fill.fore_color.rgb = RGBColor(248, 251, 253)
    bg.line.color.rgb = line; bg.line.width = Pt(1)

    kicker = slide.shapes.add_textbox(Inches(0.72), Inches(0.58), Inches(6.0), Inches(0.28))
    tf = kicker.text_frame; tf.clear(); p = tf.paragraphs[0]
    r = p.add_run(); r.text = "DECISIÓN DE COMPRA · MODELO ANALÍTICO"; r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = mid

    title = slide.shapes.add_textbox(Inches(0.72), Inches(0.86), Inches(7.8), Inches(0.6))
    tf = title.text_frame; tf.clear(); p = tf.paragraphs[0]
    r = p.add_run(); r.text = TITLE; r.font.size = Pt(27); r.font.bold = True; r.font.color.rgb = navy

    subtitle = slide.shapes.add_textbox(Inches(0.74), Inches(1.42), Inches(8.2), Inches(0.42))
    tf = subtitle.text_frame; tf.clear(); p = tf.paragraphs[0]
    r = p.add_run(); r.text = "Estima qué aparece primero, qué entra después y qué termina cerrando la compra."; r.font.size = Pt(13); r.font.color.rgb = mid

    # Stage cards
    cards = [
        (0.78, 2.12, "1", "Primero", "Qué criterio abre la decisión\nEj. tono/color"),
        (4.52, 2.12, "2", "Después", "Qué criterio entra después\nEj. precio o beneficios"),
        (8.26, 2.12, "3", "Cierre", "Qué termina definiendo\nEj. disponibilidad"),
    ]
    for i, (x, y, n, head, body) in enumerate(cards):
        card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(3.05), Inches(1.55))
        card.fill.solid(); card.fill.fore_color.rgb = pale if i == 0 else white
        card.line.color.rgb = blue if i == 0 else line; card.line.width = Pt(1.4)
        badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.18), Inches(y + 0.24), Inches(0.42), Inches(0.42))
        badge.fill.solid(); badge.fill.fore_color.rgb = blue
        badge.line.color.rgb = blue
        bt = badge.text_frame; bt.clear(); bp = bt.paragraphs[0]; bp.alignment = PP_ALIGN.CENTER
        br = bp.add_run(); br.text = n; br.font.size = Pt(14); br.font.bold = True; br.font.color.rgb = white
        box = slide.shapes.add_textbox(Inches(x + 0.75), Inches(y + 0.26), Inches(2.0), Inches(1.0))
        tf = box.text_frame; tf.clear(); p = tf.paragraphs[0]
        r = p.add_run(); r.text = head; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = navy
        p2 = tf.add_paragraph(); p2.space_before = Pt(6)
        r2 = p2.add_run(); r2.text = body; r2.font.size = Pt(11.5); r2.font.color.rgb = mid

    for x in [3.95, 7.69]:
        arr = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, Inches(x), Inches(2.69), Inches(0.44), Inches(0.34))
        arr.fill.solid(); arr.fill.fore_color.rgb = RGBColor(148, 182, 213)
        arr.line.color.rgb = RGBColor(148, 182, 213)

    # Bottom cards
    bottom = [
        (0.78, "Qué significan los %", "Probabilidades acumuladas desde el total.\nBajan conforme avanza la ruta.", "P(ruta)=P(D1)×P(D2|D1)×P(D3|D1,D2)"),
        (4.52, "Certeza", "Combina simulaciones posteriores, intervalo creíble y base efectiva.", "Alta · Media · Exploratoria"),
        (8.26, "Para qué sirve", "Traduce la decisión en acciones de anaquel, comunicación y activación.", "Ordenar · Comunicar · Cerrar"),
    ]
    for x, head, body, tag in bottom:
        card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(4.28), Inches(3.05), Inches(1.55))
        card.fill.solid(); card.fill.fore_color.rgb = white
        card.line.color.rgb = line; card.line.width = Pt(1)
        box = slide.shapes.add_textbox(Inches(x + 0.22), Inches(4.48), Inches(2.62), Inches(0.86))
        tf = box.text_frame; tf.clear(); p = tf.paragraphs[0]
        r = p.add_run(); r.text = head.upper(); r.font.size = Pt(9.5); r.font.bold = True; r.font.color.rgb = mid
        p2 = tf.add_paragraph(); p2.space_before = Pt(5)
        r2 = p2.add_run(); r2.text = body; r2.font.size = Pt(10.8); r2.font.color.rgb = navy
        pill = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x + 0.22), Inches(5.42), Inches(2.5), Inches(0.28))
        pill.fill.solid(); pill.fill.fore_color.rgb = RGBColor(234, 243, 251)
        pill.line.color.rgb = RGBColor(207, 228, 247)
        ptf = pill.text_frame; ptf.clear(); pp = ptf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
        pr = pp.add_run(); pr.text = tag; pr.font.size = Pt(8.5); pr.font.bold = True; pr.font.color.rgb = blue

    footer = slide.shapes.add_textbox(Inches(0.78), Inches(6.35), Inches(11.8), Inches(0.34))
    tf = footer.text_frame; tf.clear(); p = tf.paragraphs[0]
    r = p.add_run(); r.text = "Lectura ejecutiva: el cliente ve un Journey simple; el modelo calcula probabilidades bayesianas detrás."; r.font.size = Pt(11); r.font.color.rgb = mid

    out = io.BytesIO()
    prs.save(out)
    out.seek(0)
    return out.getvalue()
